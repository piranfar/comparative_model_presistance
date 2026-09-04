"""
Build the bioRxiv version-2 submission package from the Markdown manuscript.

Run:  python -m src.build_submission

Produces `submission/biorxiv_v2/` containing a Word manuscript with the figures
embedded at their legends, the figure files separately in PDF and PNG, and a
checklist. bioRxiv accepts DOCX or PDF; DOCX is produced here because it is the
format that survives their conversion pipeline with the fewest surprises.

The converter handles the subset of Markdown this manuscript uses: headings,
paragraphs, bold and italic runs, inline code, bullet and numbered lists,
blockquotes (used for equations), and pipe tables. It is not a general
Markdown converter and is not meant to be.
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "manuscript" / "REVISED_MANUSCRIPT.md"
OUT_DIR = ROOT / "submission" / "biorxiv_v2"
FIG_DIR = ROOT / "results" / "figures"

# figure files, in the order the legends appear
FIGURES = {
    "Figure 1.": "fig01_growth_and_strategies.png",
    "Figure 2.": "fig02_biphasic_killing.png",
    "Figure 3.": "fig03_model_fitting.png",
    "Figure 4.": "fig04_sensitivity.png",
    "Figure 5.": "fig05_mechanistic_model.png",
    "Figure 6.": "fig06_mic_mdk_plane.png",
    "Figure S1.": "fig07_supplementary_diagnostics.png",
}

INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)", re.S)


def add_runs(par, text: str) -> None:
    """Write text into a paragraph, honouring bold, italic and inline code."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        else:
            par.add_run(piece)


def add_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    header, body = cells[0], [r for r in cells[2:] if any(r)]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cs = t.add_row().cells
        for i, val in enumerate(row[:len(header)]):
            cs[i].text = ""
            add_runs(cs[i].paragraphs[0], val)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
    doc.add_paragraph()


def build() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "figures").mkdir(exist_ok=True)
    for name in FIGURES.values():
        for ext in (".png", ".pdf"):
            f = FIG_DIR / (Path(name).stem + ext)
            if f.exists():
                shutil.copy2(f, OUT_DIR / "figures" / f.name)

    lines = SRC_MD.read_text(encoding="utf-8").split("\n")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)

    i, pending_figure = 0, None
    while i < len(lines):
        ln = lines[i]

        if ln.strip() in ("---", ""):
            i += 1
            continue

        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i + 1]) <= set("|-: "):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        if ln.startswith("#"):
            level = len(ln) - len(ln.lstrip("#"))
            doc.add_heading(re.sub(r"[*`]", "", ln.lstrip("# ").strip()),
                            level=min(level, 4))
            i += 1
            continue

        if ln.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(ln.lstrip("> ").strip())
            r.font.name = "Cambria Math"
            r.italic = True
            i += 1
            continue

        if re.match(r"^\s*[-*] ", ln):
            add_runs(doc.add_paragraph(style="List Bullet"),
                     re.sub(r"^\s*[-*] ", "", ln))
            i += 1
            continue

        if re.match(r"^\s*\d+\. ", ln):
            add_runs(doc.add_paragraph(style="List Number"),
                     re.sub(r"^\s*\d+\. ", "", ln))
            i += 1
            continue

        # a figure legend: emit the image first, then the legend text
        key = next((k for k in FIGURES if ln.lstrip("*").startswith(k)), None)
        if key:
            img = OUT_DIR / "figures" / FIGURES[key]
            if img.exists():
                doc.add_picture(str(img), width=Inches(6.4))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            pending_figure = key

        par = doc.add_paragraph()
        add_runs(par, ln.strip())
        if pending_figure:
            for r in par.runs:
                r.font.size = Pt(9)
            pending_figure = None
        i += 1

    out = OUT_DIR / "Piranfar_persistence_framework_bioRxiv_v2.docx"
    doc.save(out)
    return out


def checklist() -> Path:
    text = """# bioRxiv version 2 — submission checklist

## In this folder

- `Piranfar_persistence_framework_bioRxiv_v2.docx` — the manuscript, figures embedded
- `figures/` — the seven figures as PNG (300 dpi) and PDF (vector)

## Before you upload

1. **Affiliation.** The manuscript carries a placeholder. Your three other papers
   use different affiliations, so this one has to be your call. Fill it in before
   uploading.
2. **Author contributions, competing interests, acknowledgements.** Placeholders
   in the manuscript. bioRxiv does not require them, journals do.
3. **Upload as a revision, not a new preprint.** Use the "post a revision" route
   on the existing entry, DOI 10.1101/2025.02.12.637810, so version 1 and version 2
   stay linked and readers of version 1 are shown that a correction exists.
4. **Code link.** The data and code availability section refers to the accompanying
   repository. Make that repository public first, and put its URL in that section.
   Right now it names no URL.
5. **Read the change summary once more.** It is the first thing after the title and
   it says plainly that version 1's quantitative results are withdrawn. That is
   deliberate. Anyone who cites version 1 should see it immediately.

## What is deliberately not here

The Boccarella reproduction work and the identifiability analysis of other
people's models are separate and are not part of this revision.

## Figures

| File | Shows |
|---|---|
| fig01 | growth and the three survival strategies, printed against corrected |
| fig02 | the biphasic law and its two corrections; the discontinuity |
| fig03 | model fitting on synthetic data; profile likelihood |
| fig04 | sensitivity analysis, three ways |
| fig05 | the state-structured replacement |
| fig06 | three strategies, three measurable signatures |
| fig07 | supplementary: three statistical claims checked |
"""
    p = OUT_DIR / "SUBMISSION_CHECKLIST.md"
    p.write_text(text, encoding="utf-8")
    return p


if __name__ == "__main__":
    docx = build()
    chk = checklist()
    print(f"wrote {docx}  ({docx.stat().st_size/1024:.0f} KB)")
    print(f"wrote {chk}")
    print(f"figures: {len(list((OUT_DIR / 'figures').glob('*')))} files")
